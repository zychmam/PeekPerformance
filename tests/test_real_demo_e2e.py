"""Optional end-to-end pipeline check using an actual CS2 demo file."""

from __future__ import annotations

import io
import os
import time
import zipfile
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.demos.coach_data import build_coach_inputs, load_coach_inputs
from src.demos.demo_analyzer import DemoAnalyzer
from src.demos.demo_manager import DemoManager
from src.demos.match_view import player_rounds_table, player_summary_table
from src.demos.presentation import is_steam_id_column, presentation_frame
from src.demos.rating_calculator import RATING_VERSION
from src.demos.rating_service import calculate_and_store_rating
from src.demos.versions import ANALYSIS_VERSION
from src.storage.database import Database


def _ids(frame: pd.DataFrame, column: str) -> set[str]:
    if frame.empty or column not in frame:
        return set()
    return set(frame[column].dropna().astype("string"))


@pytest.mark.skipif(
    not os.environ.get("PEEK_E2E_DEMO"), reason="Set PEEK_E2E_DEMO to a local demo file"
)
def test_real_demo_pipeline(tmp_path: Path) -> None:
    source_value = os.environ.get("PEEK_E2E_DEMO")
    if not source_value:
        pytest.skip("Set PEEK_E2E_DEMO to a local .dem, .dem.gz or .dem.zst file")
    source = Path(source_value).expanduser().resolve(strict=True)
    assert source.is_file()
    assert source.name.endswith((".dem", ".dem.gz", ".dem.zst"))

    start = time.perf_counter()
    manager = DemoManager(str(tmp_path / "demos"))
    source_bytes = source.read_bytes()
    expected_hash = manager.upload_hash(source.name, source_bytes)
    saved = manager.save_uploaded(source.name, source_bytes)
    del source_bytes
    assert saved.is_file()
    assert manager.content_hash(saved) == expected_hash

    analysis = DemoAnalyzer(saved).get_full_analysis()
    analysis.file_name = source.name.removesuffix(".gz").removesuffix(".zst")
    analysis.content_hash = manager.content_hash(saved)
    analysis.stored_file_name = saved.name
    assert analysis.header.get("map_name"), "Demo header has no map name"
    assert len(analysis.player_info) >= 2, "Parser found fewer than two players"
    assert len(analysis.rounds_df) >= 1, "Parser found no completed rounds"
    assert not analysis.kills_df.empty, "Parser found no kills"
    assert not analysis.damage_df.empty, "Parser found no damage"
    assert not analysis.round_stats_df.empty, "Round statistics are empty"
    assert not analysis.round_roster_df.empty, "Round roster is empty"
    assert analysis.analysis_version == ANALYSIS_VERSION
    if expected_hash == "c5f3f99d0cfe9eee7f1bf0dbdc2ab8ff2509a4ef01630ac3451bd9ced7dddda0":
        assert len(analysis.rounds_df) == 24  # Pre-match knife round excluded after restart.
        assert len(analysis.kills_df) == 166
        assert analysis.rounds_df.iloc[0].start_tick == 4009
    assert analysis.rounds_df["total_rounds_played"].is_monotonic_increasing
    assert analysis.damage_df["enemy_damage"].ge(0).all()
    assert analysis.damage_df["enemy_damage"].le(analysis.damage_df["actual_damage"]).all()
    for frame in [
        analysis.player_info, analysis.scoreboard, analysis.kills_df,
        analysis.damage_df, analysis.round_stats_df, analysis.grenades_df.head(100),
    ]:
        visible = presentation_frame(frame, analysis.player_info)
        assert not any(is_steam_id_column(column) for column in visible.columns)

    database = Database(str(tmp_path / "cache.db"))
    database.upsert_demo_analysis(analysis)
    coach_before = build_coach_inputs(analysis)
    coach_after = load_coach_inputs(database, analysis.content_hash)
    for source_name in coach_before.frames:
        assert coach_before.frames[source_name].event_id.tolist() == (
            coach_after.frames[source_name].event_id.tolist()
        )
    roster = database.list_demo_players(analysis.content_hash)
    assert len(roster) >= 2, "Saved demo has no player roster"
    database.track_player(roster[0]["steamid"])
    assert database.list_tracked_players()[0]["match_count"] == 1
    cached = database.get_demo_analysis(analysis.content_hash)
    assert cached is not None
    assert cached.map_name == analysis.map_name
    assert len(cached.rounds_df) == len(analysis.rounds_df)
    for frame_name, id_column in [
        ("player_info", "steamid"),
        ("kills_df", "attacker_steamid"),
        ("damage_df", "user_steamid"),
    ]:
        assert _ids(getattr(cached, frame_name), id_column) == _ids(
            getattr(analysis, frame_name), id_column
        )

    rating, per_round = calculate_and_store_rating(
        database, manager, analysis.content_hash, saved.name,
    )
    assert not rating.empty, "Rating could not be calculated"
    assert not per_round.empty
    assert rating["ADR"].ge(0).all()
    assert set(rating["steamid"].astype("string")) >= _ids(
        analysis.round_stats_df, "steamid"
    )
    player_rating = rating.set_index("steamid").loc[roster[0]["steamid"], "Rating"]
    matches = database.list_player_matches(roster[0]["steamid"], RATING_VERSION)
    assert len(matches) == 1
    assert matches[0]["rating"] == player_rating
    # A new database object simulates another app session; no demo or event reload is needed.
    reopened = Database(str(tmp_path / "cache.db"))
    with (
        patch.object(manager, "analysis_path", side_effect=AssertionError("binary reopened")),
        patch.object(reopened, "get_demo_analysis", side_effect=AssertionError("events reloaded")),
    ):
        stored_rating, stored_rounds = calculate_and_store_rating(
            reopened, manager, analysis.content_hash, saved.name,
        )
    pd.testing.assert_frame_equal(rating, stored_rating, check_dtype=False)
    pd.testing.assert_frame_equal(per_round, stored_rounds, check_dtype=False)
    totals = per_round.groupby("steamid")[["kills", "deaths", "assists", "damage"]].sum()
    for row in rating.itertuples():
        assert totals.loc[row.steamid, "kills"] == row.K
        assert totals.loc[row.steamid, "deaths"] == row.D
        assert totals.loc[row.steamid, "assists"] == row.A
        assert totals.loc[row.steamid, "damage"] / row.Rounds == pytest.approx(row.ADR, abs=0.051)

    overview = player_summary_table(cached, rating)
    assert overview.columns.tolist() == ["Gracz", "Rating", "K/D/A", "ADR", "KAST"]
    assert len(overview) >= 2
    round_view = player_rounds_table(cached, roster[0]["steamid"])
    assert len(round_view) == len(cached.rounds_df)
    assert set(round_view["Strona gracza"].dropna()).issubset({"T", "CT", "—"})

    xlsx = cached.to_excel()
    assert xlsx[:2] == b"PK"
    workbook = load_workbook(filename=io.BytesIO(xlsx), read_only=True)
    assert "Rounds" in workbook.sheetnames
    workbook.close()
    archive = io.BytesIO(cached.to_csv_zip())
    with zipfile.ZipFile(archive) as package:
        assert {"Rounds.csv", "Kills.csv", "Damage.csv"}.issubset(package.namelist())
        for name in ["Kills.csv", "Damage.csv", "PlayerInfo.csv", "RoundStats.csv"]:
            header = package.read(name).splitlines()[0].decode("utf-8")
            assert "steamid" not in header.lower().replace("_", "")
    print(
        f"E2E: map={analysis.map_name}, players={len(analysis.player_info)}, "
        f"rounds={len(analysis.rounds_df)}, kills={len(analysis.kills_df)}, "
        f"rating_rows={len(rating)}, elapsed_seconds={time.perf_counter() - start:.1f}"
    )
